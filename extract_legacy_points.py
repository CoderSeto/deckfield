#!/usr/bin/env python3
"""
Extract Legacy Points into legacy_points.json (SEASON10_PLAN B7).

    python3 extract_legacy_points.py ["Baccer Game S9 Stats.xlsm"]

Seasons 1-8 come from the workbook's `LegacyP` sheet -- teams in B2:B58,
every cell typed in (C is just SUM(E:AN)). Its column groups are:

    E-L   S1-S8 Finals   (World Championship: 16 / 8 / 4 / 2)
    N-S   S3-S8 WCS/Swiss (Cup: champion 8, runner-up 4)
    U-Z   S3-S8 M/U, LG   (League: 8 / 6 / 4 / 2; S6's two UTL champions 3 + 3)
    AB-AI S1-S8 CC        (Conference champion: 4)

Two corrections, both per explicit instruction:
- the S8 Swiss runner-up is Camphrier Town, not National Park (the HOF and
  HoFOld sheets both say Camphrier Town), so its 4 points move;
- the sheet's `Resort Area` is Battle Zone.

Season 9 is NOT read from the workbook: it is derived from the frozen record
(archive/s9_final.json's HOF_DATA and RDS_MUTUAL_STAGE), using the same
values, with the RDS cups new from S9 (winner 6, runner-up 3) and the
Regional Tournament champion standing in for the conference champion (4).
The S9 column is then checked against the totals the plan recorded.

Static history, re-runnable, like extract_hall_of_fame.py. A later season's
column is appended when that season ends (B13), never live.
"""
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "legacy_points.json")
RENAMED = {"Resort Area": "Battle Zone"}

WORLD = {16: "World Champion", 8: "World Finalist", 4: "World semifinalist", 2: "World quarterfinalist"}
LEAGUE_BY_POINTS = {8: "League champion", 6: "League second", 4: "League third", 2: "League fourth",
                    3: "UTL co-champion"}
RDS_CUPS = {"Ribbon": ("Indigo", "Silver", "Delta", "LilyValley"),
            "Dream": ("Vertress", "Phoenix", "Lanakila"),
            "Star": ("Kalosite", "Dynamax", "Terastal")}
REGION_DISPLAY = {"LilyValley": "Lily Valley"}


def col(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def plaque_label(hof, plaque, season, name):
    entry = hof["plaques"][plaque]["seasons"].get(str(season))
    if not entry:
        return None
    for label, team in zip(entry["labels"], entry["teams"]):
        if team and team["name"].rstrip("*") == name:
            return label
    return None


def region_champion_of(hof, season, name):
    for region, seasons in hof["regions"].items():
        team = seasons.get(str(season))
        if team and team["name"].rstrip("*") == name:
            return REGION_DISPLAY.get(team.get("region") or region, team.get("region") or region)
    return None


def main():
    workbook = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "Baccer Game S9 Stats.xlsm")
    wb = openpyxl.load_workbook(workbook, data_only=False, keep_vba=True)
    ws = wb["LegacyP"]
    with open(os.path.join(HERE, "hall_of_fame.json")) as f:
        hof = json.load(f)

    groups = []   # (column, season, family)
    groups += [(col("E") + i, i + 1, "World") for i in range(8)]
    groups += [(col("N") + i, i + 3, "Cup") for i in range(6)]
    groups += [(col("U") + i, i + 3, "League") for i in range(6)]
    groups += [(col("AB") + i, i + 1, "Conference") for i in range(8)]

    teams = {}

    def award(name, season, family, points, label):
        cell = teams.setdefault(name, {}).setdefault(str(season), {"points": 0, "items": []})
        cell["points"] += points
        cell["items"].append({"family": family, "points": points, "label": label})

    for r in range(2, 59):
        raw = ws.cell(r, col("B")).value
        if not raw:
            continue
        name = RENAMED.get(raw.strip(), raw.strip())
        for c, season, family in groups:
            v = ws.cell(r, c).value
            if v in (None, "", " "):
                continue
            pts = int(v)
            if family == "World":
                label = WORLD[pts]
            elif family == "Cup":
                label = plaque_label(hof, "cup", season, name) or ("Cup champion" if pts == 8 else "Cup runner-up")
            elif family == "League":
                label = plaque_label(hof, "league", season, name) or LEAGUE_BY_POINTS[pts]
            else:
                region = region_champion_of(hof, season, name)
                label = f"{region} champion" if region else "Conference champion"
            award(name, season, family, pts, label)

    # Correction: the S8 Swiss runner-up is Camphrier Town, not National Park.
    np_s8 = teams.get("National Park", {}).get("8")
    moved = [i for i in (np_s8 or {}).get("items", []) if i["family"] == "Cup" and i["points"] == 4]
    if len(moved) != 1:
        sys.exit(f"expected one S8 Cup runner-up entry on National Park, found {len(moved)}")
    np_s8["items"].remove(moved[0])
    np_s8["points"] -= 4
    if not np_s8["items"]:
        del teams["National Park"]["8"]
    award("Camphrier Town", 8, "Cup", 4, "Swiss Runner-up")

    # Season 9, from the frozen record.
    with open(os.path.join(HERE, "archive", "s9_final.json")) as f:
        s9 = json.load(f)["data"]
    h9 = s9["HOF_DATA"]
    plaque = lambda k: [t["name"] for t in h9["plaques"][k]["seasons"]["9"]["teams"]]
    champ, finalist = plaque("world")
    award(champ, 9, "World", 16, "World Champion")
    award(finalist, 9, "World", 8, "World Finalist")
    placings = [t["name"] for t in h9["placings"]["9"]]
    for n in placings[:2]:
        award(n, 9, "World", 4, "World semifinalist")
    for n in placings[2:6]:
        award(n, 9, "World", 2, "World quarterfinalist")
    pa_champ, pa_runner = plaque("cup")
    award(pa_champ, 9, "Cup", 8, "PA Cup Champion")
    award(pa_runner, 9, "Cup", 4, "PA Cup Runner-up")
    for cup, winner in zip(("Ribbon", "Dream", "Star"), plaque("rds")):
        final = s9["RDS_MUTUAL_STAGE"][cup]["final"]
        runner = final["away"] if final["home"] == winner else final["home"]
        award(winner, 9, "RDS", 6, f"{cup} Cup Winner")
        award(runner, 9, "RDS", 3, f"{cup} Cup Runner-up")
    for pts, (label, name) in zip((8, 6, 4, 2), zip(h9["plaques"]["league"]["seasons"]["9"]["labels"],
                                                    plaque("league"))):
        award(name, 9, "League", pts, label)
    for region, team in h9["regions"].items():
        if team.get("9"):
            award(team["9"]["name"], 9, "Conference",
                  4, f"{REGION_DISPLAY.get(region, region)} RT champion")

    # The S9 column must match what the plan recorded.
    s9_totals = {n: v["9"]["points"] for n, v in teams.items() if "9" in v}
    expected = {"Casseroya Lake": 26, "Canalave City": 16, "Snowpoint City": 13, "Nimbasa City": 11,
                "Camphrier Town": 8, "Petalburg City": 8, "Cabo Poco": 8, "National Park": 8,
                "Pueltown": 6, "Cocona Village": 6, "Mesagoza": 5, "Victory Road": 4,
                "Mistralton City": 4, "Po Town": 4, "Stow-on-Side": 4, "Dolce Island": 4,
                "Blueberry Terarium": 2, "Vermilion City": 2}
    if s9_totals != expected:
        sys.exit(f"S9 legacy column disagrees with SEASON10_PLAN B7: {s9_totals}")
    totals = {n: sum(c["points"] for c in v.values()) for n, v in teams.items()}
    if len(totals) != 67 or max(totals.items(), key=lambda kv: kv[1]) != ("Canalave City", 70):
        sys.exit(f"expected 67 teams with Canalave City leading on 70, got {len(totals)} and "
                 f"{max(totals.items(), key=lambda kv: kv[1])}")

    out = {
        "note": ("Legacy Points (SEASON10_PLAN B7). Seasons 1-8 from the workbook's LegacyP sheet "
                 "(S8 Swiss runner-up corrected to Camphrier Town; Resort Area is Battle Zone); "
                 "Season 9 from archive/s9_final.json. A season's column is added when it ends."),
        "values": {"World": "Champion 16, Finalist 8, semifinalists 4, quarterfinalists 2",
                   "Cup": "WCS / Swiss / PA Cup: champion 8, runner-up 4",
                   "RDS": "Ribbon / Dream / Star (from S9): winner 6, runner-up 3",
                   "League": "8 / 6 / 4 / 2 (S6's two UTL champions 3 + 3)",
                   "Conference": "Conference (from S9, Regional Tournament) champion 4"},
        "seasons": list(range(1, 10)),
        "teams": {n: {s: teams[n][s] for s in sorted(teams[n], key=int)}
                  for n in sorted(teams, key=lambda n: (-totals[n], n))},
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"Wrote {os.path.relpath(OUT, HERE)}: {len(totals)} teams; "
          f"S9 column {sum(s9_totals.values())} points across {len(s9_totals)} teams; "
          f"Canalave City leads all-time with {totals['Canalave City']}")


if __name__ == "__main__":
    main()
