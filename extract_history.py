"""Extract a finished pre-S9 season from its workbook into history/s<N>/.

    python3 extract_history.py 8 "history/s8/Baccer Game S8 Stats.xlsm"

The historical archive keeps exactly what the season RECORDED. Nothing here is
simulated, estimated or recomputed with today's formulas: a stat the workbook
never recorded stays blank. Anything filled in later ("relitigating") belongs
in a separate layer beside these files, never in them.

What it writes, all under history/s<N>/:

- games.csv      one row per game, Archive order. team_a is the HOST (home is
                 always "A", as in the modern CSV). Team ids are CANONICAL --
                 the team's S9 dex -- not the season's own dex, which is not a
                 stable identity (see history/teams.json).
                 rank_a/rank_b are each side's rank GOING INTO the game (Archive
                 H/M). goals_a/goals_b are goals scored (Archive N, written
                 [winner-loser]); elo_k is the game's Elo K (Archive E): the
                 points the winner gained and the loser lost, exactly.
                 raw_winner/raw_loser
                 are the two packed game strings verbatim, for audit.
- walkovers.csv  byes the season scored as wins (S8: the Grand Swiss
                 round-1/round-2 byes, from Calcs' "3S-" cells).
- teams.csv      each team as it stood that season: its own dex, canonical id,
                 name, region, division, final rank / OVR / Elo, accolades.
- events.json    every event in play order with its game count (an event with
                 no games, like a leg 3 nobody had to play, is listed and
                 marked empty), plus the source workbook's SHA-256.

It also adds the season to history/teams.json, the cross-season identity map.

It refuses to write unless every check against the workbook's own totals and
against later records passes -- see validate(). Season-specific layout lives in
SEASONS; S8 is the only one so far.
"""
import csv
import hashlib
import json
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
HISTORY = os.path.join(ROOT, "history")

GOALS_PATTERN = re.compile(r'^\[(\d+)-(\d+)\]$')  # Archive N: [winner goals-loser goals]
GAME_PATTERN = re.compile(
    r'^(\d)([RLSPF])!(-?\d+)\+(-?\d+)-(-?\d+)\^(-?[\d.]+)\((-?[\d.]+)\)'
    r'\[(-?[\d.]+)\]\{(-?[\d.]+)\}(-?[\d.]+)$'
)
WALKOVER_PATTERN = re.compile(r'^(\d)([RLSPF])-$')
EX_PATTERN = re.compile(r'^/(-?\d+)/$')

LEAGUE_DIVISION_MAP = {
    "First": 1, "Second": 2, "Third": 3, "Fourth": 4, "Fifth": 5,
    "Sixth": 6, "Seventh": 7, "Eighth": 8, "Ninth": 9, "Tenth": 10,
}
REGION_NAME_FIX = {"Lily Valley": "LilyValley"}
NUMBER_WORDS = {w: i for i, w in enumerate(
    "ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN ELEVEN TWELVE THIRTEEN "
    "FOURTEEN FIFTEEN".split(), start=1)}

# A team renamed since that season, old name -> its S9 name. Each entry is
# proved by S9's own "S8 End" rank checkpoint, which lists the team under its
# S9 name with the rank it finished S8 on (validate() checks all 160).
RENAMED_SINCE = {
    "Resort Area": "Battle Zone",
    "Fight Area": "Cabo Poco",
}

SEASONS = {
    8: {
        # Rankings columns (0-based): A dex, B region, C name, AN rank,
        # AQ OVR, BL Elo, BQ base climate, BV division, BY primary type,
        # CA accolades, CK the FINAL rank checkpoint, FV the Elo carried in
        # from S7 (proved: it is each team's Elo before its first S8 game).
        "rankings": dict(dex=0, region=1, name=2, rank=39, ovr=42, elo=63,
                         basclm=68, division=73, primary=76, accolades=78,
                         final_rank=88, opening_elo=177),
        # Rankings F..M: the W-L the workbook itself kept per bucket.
        "record_cols": dict(RW=5, RL=6, LW=7, LL=8, SW=9, SL=10, pW=11, pL=12),
        # Calcs: per-team game grid, column L onward in play order.
        "calcs_first_col": 12,
    },
}


# ---------------------------------------------------------------------------
# Events
# ---------------------------------------------------------------------------

def _is_title(label):
    """An event title is ALL-CAPS ("GRAND SWISS FINAL - First Leg" starts with
    one); a sub-label ("First Leg", "The Overhaul") refines the current one."""
    first = label.split()[0]
    return first.isupper() and len(first) > 1


def _event_code(title, sub):
    """A short, stable code for an S8 event, from its Archive labels."""
    t = title.strip().upper()
    s = (sub or "").strip()
    leg = re.search(r'(First|Second) Leg', s + " " + title, re.I)
    leg_n = {"FIRST": 1, "SECOND": 2}[leg.group(1).upper()] if leg else None
    game = re.search(r'Game (One|Two|Three)', s + " " + title, re.I)
    game_n = NUMBER_WORDS[game.group(1).upper()] if game else None

    m = re.match(r'(REGIONAL|LEAGUE) MATCHDAY (\w+)$', t)
    if m:
        return ("R" if m.group(1) == "REGIONAL" else "L") + str(NUMBER_WORDS[m.group(2)])
    m = re.match(r'GRAND SWISS ROUND (\w+)$', t)
    if m and m.group(1) in NUMBER_WORDS:
        return f"GS{NUMBER_WORDS[m.group(1)]}"
    stage = {
        "GRAND SWISS PLAY-IN ROUND": "GS-PI", "GRAND SWISS ROUND OF SIXTEEN": "GS-R16",
        "GRAND SWISS QUARTERFINAL": "GS-QF", "GRAND SWISS SEMIFINAL": "GS-SF",
        "CHAMPIONSHIP PLAY-IN": "WC-PI",
    }.get(t)
    if stage:
        return f"{stage}-{leg_n}"
    if t.startswith("GRAND SWISS FINAL"):
        return f"GS-F-{leg_n}"
    if t == "REGIONAL PLAYOFFS":
        rnd = {"Quarterfinal": "QF", "Semifinal": "SF", "Championship": "F"}[s.split(" - ")[0]]
        return f"RP-{rnd}-{leg_n}"
    if t == "CHAMPIONSHIP SWISS":
        return f"WC-S{NUMBER_WORDS[s.split()[-1].upper()]}"
    if t == "CHAMPIONSHIP FINALS":
        rnd = {"Play-in": "FPI", "Round of Sixteen": "R16", "Quarterfinal": "QF",
               "Semifinal": "SF"}[s.split(" - ")[0]]
        return f"WC-{rnd}-{leg_n or game_n}"
    if t.startswith("WORLD CHAMPIONSHIP - GAME"):
        return f"WC-F-{game_n}"
    raise ValueError(f"unrecognised event: {title!r} / {sub!r}")


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------

def read_teams(wb, cfg):
    c = cfg["rankings"]
    teams = {}
    for r in wb["Rankings"].iter_rows(min_row=2, max_row=161, values_only=True):
        if r[c["dex"]] is None:
            continue
        region = r[c["region"]]
        teams[r[c["dex"]]] = {
            "dex": r[c["dex"]], "name": r[c["name"]],
            "region": REGION_NAME_FIX.get(region, region),
            "division": LEAGUE_DIVISION_MAP[r[c["division"]]],
            "primary_type": r[c["primary"]], "basclm": r[c["basclm"]],
            "final_rank": r[c["final_rank"]], "final_ovr": r[c["ovr"]],
            "final_elo": r[c["elo"]], "opening_elo": r[c["opening_elo"]],
            "accolades": r[c["accolades"]] or "",
            "record": {k: r[i] for k, i in cfg["record_cols"].items()},
        }
    return teams


def read_games(wb):
    """Every Archive game, in sheet order, with its event."""
    rows = list(wb["Archive"].iter_rows(min_row=1, max_col=23, values_only=True))
    events, games = [], []
    cur = None
    for rownum, r in enumerate(rows, start=1):
        if rownum == 1:
            continue
        label = r[1]
        if isinstance(label, str) and label.strip():
            if _is_title(label):
                cur = {"title": label.strip(), "sub": None, "row": rownum, "games": 0}
                events.append(cur)
            elif cur is not None and cur["sub"] is None:
                cur["sub"] = label.strip()
        if not r[15]:
            continue
        w, l = GAME_PATTERN.match(str(r[15])), GAME_PATTERN.match(str(r[16]))
        if not w or not l:
            raise ValueError(f"Archive row {rownum}: unparseable game {r[15]!r} / {r[16]!r}")
        cur["games"] += 1
        ex = None
        if r[18] not in (None, ""):
            m = EX_PATTERN.match(str(r[18]).strip())
            if not m:
                raise ValueError(f"Archive row {rownum}: unparseable EX {r[18]!r}")
            ex = int(m.group(1))
        games.append({
            "row": rownum, "event": cur, "winner_name": r[8], "loser_name": r[11],
            "home_away": str(r[6]).strip().upper(), "w": w.groups(), "l": l.groups(),
            "ex": ex, "mark": r[17] or "", "note": r[19] or "",
            "rank_w": r[7], "rank_l": r[12], "elo_k": r[4], "goals": r[13],
        })
    for i, e in enumerate(events, start=1):
        e["round"] = i
        e["code"] = _event_code(e["title"], e["sub"])
    return events, games


def read_walkovers(wb, cfg, events):
    """Calcs' "3S-" cells: byes scored as wins. The Calcs column's header names
    the slot (S1, S2...), which is matched to that Grand Swiss round's event."""
    ws = wb["Calcs"]
    first = cfg["calcs_first_col"]
    header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    by_code = {e["code"]: e for e in events}
    out = []
    for r in ws.iter_rows(min_row=3, max_row=162, values_only=True):
        if r[0] is None:
            continue
        for j in range(first - 1, len(r)):
            v = r[j]
            if not isinstance(v, str):
                continue
            m = WALKOVER_PATTERN.match(v)
            if not m:
                continue
            slot = header[j]
            sm = re.match(r'^S(\d+)$', str(slot))
            if not sm or f"GS{sm.group(1)}" not in by_code:
                raise ValueError(f"walkover {v!r} for dex {r[0]} in unexpected Calcs slot {slot!r}")
            ev = by_code[f"GS{sm.group(1)}"]
            out.append({"dex": r[0], "result": int(m.group(1)), "game_type": m.group(2),
                        "round": ev["round"], "event": ev["code"]})
    return out


# ---------------------------------------------------------------------------
# Identity: the season's dex -> canonical (S9) id
# ---------------------------------------------------------------------------

def canonical_ids(teams):
    s9 = json.load(open(os.path.join(ROOT, "archive", "s9_inputs.json")))["teams"]
    by_name = {t["name"]: int(dex) for dex, t in s9.items()}
    out = {}
    for dex, t in teams.items():
        name = RENAMED_SINCE.get(t["name"], t["name"])
        if name not in by_name:
            raise ValueError(f"{t['name']!r} (dex {dex}) has no S9 counterpart; add it to RENAMED_SINCE")
        out[dex] = by_name[name]
    if sorted(out.values()) != sorted(by_name.values()):
        raise ValueError("the season's teams do not map one-to-one onto S9's")
    return out, {cid: s9[str(cid)]["name"] for cid in out.values()}


# ---------------------------------------------------------------------------
# Building rows and validating
# ---------------------------------------------------------------------------

def build_game_rows(games, name_to_dex, canon):
    rows = []
    for seq, g in enumerate(games, start=1):
        w, l = g["w"], g["l"]
        wdex, ldex = name_to_dex[g["winner_name"]], name_to_dex[g["loser_name"]]
        winner_home = g["home_away"] == "H"
        home, away = (w, l) if winner_home else (l, w)
        rank_h, rank_a = (g["rank_w"], g["rank_l"]) if winner_home else (g["rank_l"], g["rank_w"])
        gw, gl = (int(x) for x in GOALS_PATTERN.match(str(g["goals"])).groups())
        goals_h, goals_a = (gw, gl) if winner_home else (gl, gw)
        hdex, adex = (wdex, ldex) if winner_home else (ldex, wdex)
        ot = "OT" in g["note"] or (int(w[0]) == 2)
        rows.append({
            "seq": seq, "round": g["event"]["round"], "event": g["event"]["code"],
            "game_type": w[1],
            "team_a": canon[hdex], "team_b": canon[adex],
            "result_a": int(home[0]), "pf_a": int(home[3]), "pa_a": int(home[4]),
            "raw_goal_a": home[5], "raw_goal_b": away[5],
            "spread_a": home[6], "interest": home[7],
            "dscr_a": home[8], "dscr_b": away[8],
            "elo_a_after": home[9], "elo_b_after": away[9],
            "home": "A",
            # Modern per-team EX is a gap: this season banked a bonus for the
            # winner only, and as a differential, not each side's own score.
            "ex_a": "", "ex_b": "",
            "ex_winner_diff": "" if g["ex"] is None else g["ex"],
            "winner": "A" if winner_home else "B",
            "goals_a": goals_h, "goals_b": goals_a, "elo_k": g["elo_k"],
            "rank_a": rank_h, "rank_b": rank_a,
            "ot": int(ot), "mark": g["mark"], "note": g["note"].replace("\n", " "),
            "source_row": g["row"], "raw_winner": "".join(_raw(w)), "raw_loser": "".join(_raw(l)),
        })
    return rows


def _raw(g):
    r, t, opp, pf, pa, raw, spread, gi, dscr, elo = g
    return [f"{r}{t}!{opp}+{pf}-{pa}^{raw}({spread})[{gi}]{{{dscr}}}{elo}"]


GAME_FIELDS = ["seq", "round", "event", "game_type", "team_a", "team_b", "result_a",
               "pf_a", "pa_a", "raw_goal_a", "raw_goal_b", "spread_a", "interest",
               "dscr_a", "dscr_b", "elo_a_after", "elo_b_after", "home", "ex_a", "ex_b",
               "ex_winner_diff", "winner", "goals_a", "goals_b", "elo_k", "rank_a", "rank_b",
               "ot", "mark", "note", "source_row", "raw_winner", "raw_loser"]


def validate(season, teams, games, walkovers, canon, events):
    """Every check must pass or nothing is written."""
    errors = []
    name_to_dex = {t["name"]: d for d, t in teams.items()}

    # 1. Both sides of every game agree, and the strings name the right teams.
    for g in games:
        w, l = g["w"], g["l"]
        where = f"row {g['row']}"
        if int(w[0]) + int(l[0]) != 3:
            errors.append(f"{where}: results {w[0]}/{l[0]} do not sum to 3")
        if (w[3], w[4]) != (l[4], l[3]):
            errors.append(f"{where}: scores do not mirror")
        if abs(float(w[6]) + float(l[6])) > 0.011:
            errors.append(f"{where}: spreads do not mirror")
        if w[7] != l[7] or w[1] != l[1]:
            errors.append(f"{where}: interest or game type differ")
        if int(w[0]) < 2:
            errors.append(f"{where}: the named winner did not win")
        if int(w[2]) != name_to_dex.get(g["loser_name"]) or int(l[2]) != name_to_dex.get(g["winner_name"]):
            errors.append(f"{where}: opponent dex does not match the named teams")
        if g["home_away"] not in ("H", "A"):
            errors.append(f"{where}: host {g['home_away']!r} is neither H nor A")
        gm = GOALS_PATTERN.match(str(g["goals"]))
        if not gm:
            errors.append(f"{where}: goals {g['goals']!r} unparseable")
        elif int(gm.group(1)) > int(w[3]) or int(gm.group(2)) > int(l[3]):
            errors.append(f"{where}: more goals than points (goals {g['goals']}, score {w[3]}-{l[3]})")
        if not isinstance(g["elo_k"], (int, float)):
            errors.append(f"{where}: Elo K {g['elo_k']!r} is not a number")

    # 2. Per-team W-L in each bucket equals the workbook's own, byes included.
    bucket = {"R": "R", "L": "L", "S": "S", "P": "p", "F": "p"}
    rec = defaultdict(Counter)
    for g in games:
        rec[name_to_dex[g["winner_name"]]][bucket[g["w"][1]] + "W"] += 1
        rec[name_to_dex[g["loser_name"]]][bucket[g["l"][1]] + "L"] += 1
    for wo in walkovers:
        rec[wo["dex"]][bucket[wo["game_type"]] + ("W" if wo["result"] >= 2 else "L")] += 1
    for dex, t in teams.items():
        for k, v in t["record"].items():
            if rec[dex][k] != v:
                errors.append(f"{t['name']}: {k} {rec[dex][k]} from the games, {v} in Rankings")

    # 3. The Elo chain: from each team's opening Elo, every game moves the
    #    winner up and the loser down by exactly that game's Elo K (Archive E,
    #    the points exchanged).
    elo = {d: float(t["opening_elo"]) for d, t in teams.items()}
    for g in games:
        wd, ld = name_to_dex[g["winner_name"]], name_to_dex[g["loser_name"]]
        k = float(g["elo_k"]) if isinstance(g["elo_k"], (int, float)) else None
        if k is not None and (abs(elo[wd] + k - float(g["w"][9])) > 0.5
                              or abs(elo[ld] - k - float(g["l"][9])) > 0.5):
            errors.append(f"row {g['row']}: Elo {elo[wd]:.0f}/{elo[ld]:.0f} +/- K {k} != "
                          f"{g['w'][9]}/{g['l'][9]}")
        elo[wd], elo[ld] = float(g["w"][9]), float(g["l"][9])

    # 3b. Each team's last Elo equals its final Elo.
    last = {}
    for g in games:
        last[name_to_dex[g["winner_name"]]] = float(g["w"][9])
        last[name_to_dex[g["loser_name"]]] = float(g["l"][9])
    for dex, t in teams.items():
        if last.get(dex) != float(t["final_elo"]):
            errors.append(f"{t['name']}: last game Elo {last.get(dex)} != final {t['final_elo']}")

    # 4. The next season's own record of how this one ended: S9's "S8 End"
    #    ranks (copied verbatim from S9's workbook) by canonical name. This is
    #    what proves RENAMED_SINCE.
    if season == 8:
        s8_end = json.load(open(os.path.join(ROOT, "seasons", "s9", "rank_history_verbatim.json")))["S8 End"]
        s9_names = json.load(open(os.path.join(ROOT, "archive", "s9_inputs.json")))["teams"]
        for dex, t in teams.items():
            nm = s9_names[str(canon[dex])]["name"]
            if s8_end.get(nm) != t["final_rank"]:
                errors.append(f"{t['name']} -> {nm}: final rank {t['final_rank']}, S9's S8 End says {s8_end.get(nm)}")

    # 5. The Hall of Fame's titles for this season follow from the games.
    errors += _check_hall_of_fame(season, games, teams, events)

    # 6. Regional and League are complete single round robins: each pair in a
    #    region (division) meets exactly once, never across groups, and every
    #    team plays every matchday.
    for gtype, key in (("R", "region"), ("L", "division")):
        pairs, per_event = Counter(), defaultdict(Counter)
        for g in games:
            if g["w"][1] != gtype:
                continue
            a, b = name_to_dex[g["winner_name"]], name_to_dex[g["loser_name"]]
            pairs[frozenset((a, b))] += 1
            per_event[g["event"]["code"]].update((a, b))
            if teams[a][key] != teams[b][key]:
                errors.append(f"row {g['row']}: {gtype} game across {key}s")
        sizes = Counter(t[key] for t in teams.values())
        expected = sum(n * (n - 1) // 2 for n in sizes.values())
        if len(pairs) != expected or any(v != 1 for v in pairs.values()):
            errors.append(f"{gtype}: {len(pairs)} distinct pairs (expected {expected}), "
                          f"{sum(v > 1 for v in pairs.values())} repeated")
        for code, c in per_event.items():
            if len(c) != len(teams) or set(c.values()) != {1}:
                errors.append(f"{code}: not every team played exactly once")

    return errors


def _tie_winner(games):
    """Aggregate over a two-legged tie; returns (winner, {name: points})."""
    pts = Counter()
    for g in games:
        pts[g["winner_name"]] += int(g["w"][3])
        pts[g["loser_name"]] += int(g["l"][3])
    (a, pa), (b, pb) = pts.most_common(2)
    return (a if pa != pb else None), pts


def derived_titles(games, teams, events):
    """Champions this season's games decide: the World Championship (best of
    three), the Grand Swiss (two-legged final) and each region's playoff (a
    two-legged final). An exact aggregate tie is reported, not guessed."""
    by_code = defaultdict(list)
    for g in games:
        by_code[g["event"]["code"]].append(g)
    titles = {}

    wc = Counter(g["winner_name"] for c in ("WC-F-1", "WC-F-2", "WC-F-3") for g in by_code[c])
    finalists = {g[k] for c in ("WC-F-1", "WC-F-2", "WC-F-3") for g in by_code[c]
                 for k in ("winner_name", "loser_name")}
    champ = wc.most_common(1)[0][0]
    titles["world"] = (champ, (finalists - {champ}).pop())

    gs = by_code["GS-F-1"] + by_code["GS-F-2"]
    w, pts = _tie_winner(gs)
    titles["swiss"] = (w, (set(pts) - {w}).pop() if w else None)

    region_of = {t["name"]: t["region"] for t in teams.values()}
    ties = defaultdict(list)
    for g in by_code["RP-F-1"] + by_code["RP-F-2"]:
        ties[region_of[g["winner_name"]]].append(g)
    titles["regions"] = {}
    for region, gl in ties.items():
        w, pts = _tie_winner(gl)
        titles["regions"][region] = (w, (set(pts) - {w}).pop() if w else None, dict(pts))
    return titles


def _check_hall_of_fame(season, games, teams, events):
    hof = json.load(open(os.path.join(ROOT, "hall_of_fame.json")))
    s = str(season)
    errors = []
    titles = derived_titles(games, teams, events)

    def hof_names(entry):
        return [RENAMED_SINCE.get(t["name"], t["name"]) if t else None for t in entry["teams"]]

    def canonical(n):
        return RENAMED_SINCE.get(n, n)

    world = hof_names(hof["plaques"]["world"]["seasons"][s])
    if [canonical(x) for x in titles["world"]] != world[:2]:
        errors.append(f"World Championship: games say {titles['world']}, Hall of Fame {world[:2]}")

    cup = hof["plaques"]["cup"]["seasons"][s]
    if [canonical(x) if x else None for x in titles["swiss"]] != hof_names(cup)[:2]:
        errors.append(f"Grand Swiss: games say {titles['swiss']}, Hall of Fame {hof_names(cup)}")

    for region, (w, ru, pts) in titles["regions"].items():
        if w is None:
            errors.append(f"{region} regional final is level on aggregate {pts}: no rule to settle it")
            continue
        plate = hof["regions"].get(region, {}).get(s)
        if plate is None:
            errors.append(f"{region}: no Hall of Fame entry for S{s}")
            continue
        champ = plate["name"]
        if canonical(w) != RENAMED_SINCE.get(champ, champ):
            errors.append(f"{region} regional champion: games say {w}, Hall of Fame {champ}")
    return errors


# ---------------------------------------------------------------------------

def write_csv(path, fields, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def main(season, workbook):
    cfg = SEASONS[season]
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    teams = read_teams(wb, cfg)
    events, games = read_games(wb)
    walkovers = read_walkovers(wb, cfg, events)
    canon, canon_names = canonical_ids(teams)

    errors = validate(season, teams, games, walkovers, canon, events)
    if errors:
        print(f"REFUSING to write S{season}: {len(errors)} check(s) failed")
        for e in errors[:50]:
            print("  " + e)
        sys.exit(1)

    out = os.path.join(HISTORY, f"s{season}")
    os.makedirs(out, exist_ok=True)
    name_to_dex = {t["name"]: d for d, t in teams.items()}

    write_csv(os.path.join(out, "games.csv"), GAME_FIELDS, build_game_rows(games, name_to_dex, canon))
    write_csv(os.path.join(out, "walkovers.csv"),
              ["round", "event", "team", "game_type", "result"],
              [{"round": w["round"], "event": w["event"], "team": canon[w["dex"]],
                "game_type": w["game_type"], "result": w["result"]}
               for w in sorted(walkovers, key=lambda w: (w["round"], canon[w["dex"]]))])
    team_fields = ["team", "season_dex", "name", "region", "division", "primary_type",
                   "basclm", "opening_elo", "final_rank", "final_ovr", "final_elo", "accolades"]
    write_csv(os.path.join(out, "teams.csv"), team_fields, [
        {"team": canon[d], "season_dex": d, "name": t["name"], "region": t["region"],
         "division": t["division"], "primary_type": t["primary_type"], "basclm": t["basclm"],
         "opening_elo": t["opening_elo"], "final_rank": t["final_rank"], "final_ovr": t["final_ovr"], "final_elo": t["final_elo"],
         "accolades": t["accolades"].replace("\n", " ")}
        for d, t in sorted(teams.items(), key=lambda kv: canon[kv[0]])])

    sha = hashlib.sha256(open(workbook, "rb").read()).hexdigest()
    titles = derived_titles(games, teams, events)
    with open(os.path.join(out, "events.json"), "w") as f:
        json.dump({
            "season": season,
            "source": {"workbook": os.path.relpath(workbook, ROOT), "sha256": sha},
            "games": len(games), "walkovers": len(walkovers),
            "titles": {"world": titles["world"], "grand_swiss": titles["swiss"],
                       "regions": {r: v[:2] for r, v in sorted(titles["regions"].items())}},
            "events": [{"round": e["round"], "code": e["code"], "title": e["title"],
                        "sub": e["sub"], "games": e["games"], "empty": e["games"] == 0,
                        "source_row": e["row"]} for e in events],
        }, f, indent=1)
        f.write("\n")

    # The cross-season identity map.
    tpath = os.path.join(HISTORY, "teams.json")
    ident = json.load(open(tpath)) if os.path.exists(tpath) else {}
    for d, t in teams.items():
        cid = str(canon[d])
        entry = ident.setdefault(cid, {"s9_name": canon_names[canon[d]], "seasons": {}})
        entry["seasons"][str(season)] = {"dex": d, "name": t["name"], "region": t["region"],
                                         "division": t["division"]}
    ident = {k: ident[k] for k in sorted(ident, key=int)}
    with open(tpath, "w") as f:
        json.dump(ident, f, indent=1)
        f.write("\n")

    print(f"S{season}: {len(games)} games, {len(walkovers)} walkovers, {len(events)} events "
          f"({sum(e['games'] == 0 for e in events)} empty), {len(teams)} teams -> {os.path.relpath(out, ROOT)}/")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(int(sys.argv[1]), sys.argv[2])
